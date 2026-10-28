"""Shared, explicit input normalization and validated non-secret study data.

No model requests, prompt assembly, or editor mutation happen in this module.
The same entry points are used by file imports and, in later stages, UI/CLI.
"""
from __future__ import annotations

import copy
import csv
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from functools import lru_cache
import hashlib
from importlib import resources
import io
import math
import os
from pathlib import Path
import re
from types import MappingProxyType
from typing import Any, Literal, Mapping
from urllib.parse import parse_qsl, urlsplit
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError, fromstring

from pydantic import BaseModel, ConfigDict, Field, SecretStr, ValidationError, field_serializer, model_validator
import yaml

BUILTIN_DATASET_ID = "anes2016"
BUILTIN_DATASET_FILE = "anes2016_personas.csv"
BUILTIN_DATASET_ROWS = 2708
BUILTIN_DATASET_SHA256 = "ab0064577dc0afd8c1ffa9946ca541318f1fd735d3c0358d40364350909cb7f7"
DEMOGRAPHIC_FIELDS = ("ideology", "age", "race_ethnicity", "gender", "education", "party")
BIG_FIVE_FIELDS = ("openness", "conscientiousness", "extraversion", "agreeableness", "neuroticism")
PERSONA_FIELDS = DEMOGRAPHIC_FIELDS + BIG_FIVE_FIELDS
PERSONA_COLUMNS = ("persona_id",) + PERSONA_FIELDS
ITEM_COLUMNS = ("item_id", "scale", "text", "type", "min", "max", "min_label", "max_label", "pre", "turn", "post")
PHASES = ("pre", "turn", "post")
SURVEY_INSTRUCTIONS = {
    "pre": "Before the conversation begins, please answer the following questions as yourself.",
    "turn": "The conversation is in progress. Please answer based on your experience and thoughts at the current moment, considering the conversation so far.",
    "post": "The conversation has ended. Thinking about the whole conversation, please answer the following questions as yourself.",
}
# prompts.py owns registry defaults, wording, and placeholder validation.
TEMPLATE_KEYS = frozenset((
    "system", "main_user", "survey_user", "private_state_block", "state_guidance",
    "state_init_user", "state_update_user", "state_response_instruction",
    "state_transcript_message", "state_transcript_silence", "main_state_guidance",
    "survey_state_guidance", "survey_guidance", "pass_line", "early_end_line",
    "reaction_line", "transcript_separator", "transcript_message",
    "survey_transcript_message", "transcript_silence", "survey_transcript_silence",
    "silence_guidance", "main_empty_context", "empty_public_transcript", "pre_context",
    "turn_context", "post_context", "likert_question", "text_question", "question_separator",
    "big_five_line", "big_five_trait", "big_five_separator", "pre_survey_block",
    "pre_memory_likert_answer", "pre_memory_text_answer", "pre_memory_separator",
))


class ConfigError(ValueError):
    """A plain-language validation error safe to display without source values."""


class DatasetError(ConfigError):
    """An unavailable or damaged bundled dataset; no fallback is permitted."""


def is_blank(value: Any) -> bool:
    # In particular NaN is NOT a blank cell; numerical NaN must be rejected.
    return value is None or (isinstance(value, str) and not value.strip())


def normalize_text(value: Any, field: str, *, required: bool = True, trim: bool = True) -> str | None:
    if is_blank(value):
        if required:
            raise ConfigError(f"{field}: enter non-empty text.")
        return None
    if not isinstance(value, str):
        raise ConfigError(f"{field}: enter text, not a number or another value type.")
    return value.strip() if trim else value


def normalize_number(value: Any, field: str, *, integer: bool = False, minimum: float | None = None,
                     maximum: float | None = None, optional: bool = False) -> int | float | None:
    if optional and is_blank(value):
        return None
    if isinstance(value, bool) or isinstance(value, (dict, list, tuple)) or is_blank(value):
        raise ConfigError(f"{field}: enter a finite {'whole number' if integer else 'number'}; booleans are not numbers.")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise ConfigError(f"{field}: enter a finite {'whole number' if integer else 'number'}.") from None
    if not number.is_finite():
        raise ConfigError(f"{field}: NaN and infinity are not allowed.")
    if integer and number != number.to_integral_value():
        raise ConfigError(f"{field}: enter a whole number (for example 34 or 34.0).")
    if minimum is not None and number < Decimal(str(minimum)):
        raise ConfigError(f"{field}: must be at least {minimum:g}.")
    if maximum is not None and number > Decimal(str(maximum)):
        raise ConfigError(f"{field}: must be at most {maximum:g}.")
    result = int(number) if integer else float(number)
    if not integer and not math.isfinite(result):
        raise ConfigError(f"{field}: the number is too large to represent as a finite value.")
    return result


def normalize_boolean(value: Any, field: str, *, table: bool = False) -> bool:
    if type(value) is bool:
        return value
    if table and isinstance(value, str) and value.strip().lower() in ("true", "false"):
        return value.strip().lower() == "true"
    raise ConfigError(f"{field}: use {'TRUE or FALSE' if table else 'an actual true or false boolean'}, not 1/0 or other text.")


def normalize_id(value: Any, field: str, *, optional: bool = False, yaml_input: bool = False) -> str | None:
    if optional and is_blank(value):
        return None
    if yaml_input and not isinstance(value, str):
        raise ConfigError(f'{field}: must be a string in YAML; write it in quotes, for example persona_id: "2".')
    if not isinstance(value, str):
        # Excel numeric cells have already lost their formatting; never invent zeros.
        number = normalize_number(value, field)
        value = str(int(number)) if number.is_integer() else str(number)
    return normalize_text(value, field, required=not optional)


def _mapping(value: Any, field: str) -> dict:
    if not isinstance(value, Mapping):
        raise ConfigError(f"{field}: expected an object with named fields.")
    if any(not isinstance(key, str) for key in value):
        raise ConfigError(f"{field}: field names must be strings.")
    return dict(value)


def _persona_fields(data: dict, *, complete: bool) -> dict:
    result = dict(data)
    for name in DEMOGRAPHIC_FIELDS:
        if name in result or complete:
            if name == "age":
                result[name] = normalize_number(result.get(name), name, integer=True, optional=not complete)
            else:
                result[name] = normalize_text(result.get(name), name, required=complete)
    for name in BIG_FIVE_FIELDS:
        if name in result:
            result[name] = normalize_number(result[name], name, minimum=1, maximum=5, optional=True)
    if "persona_id" in result:
        result["persona_id"] = normalize_id(result["persona_id"], "persona_id", optional=True)
    return result


class FrozenModel(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, frozen=True, hide_input_in_errors=True)


class Persona(FrozenModel):
    persona_id: str | None = None
    ideology: str
    age: int
    race_ethnicity: str
    gender: str
    education: str
    party: str
    openness: float | None = None
    conscientiousness: float | None = None
    extraversion: float | None = None
    agreeableness: float | None = None
    neuroticism: float | None = None

    @model_validator(mode="before")
    @classmethod
    def normalize(cls, value):
        return _persona_fields(_mapping(value, "persona"), complete=True)


class AgentConfig(FrozenModel):
    display_name: str | None = None
    model: str
    api_key: SecretStr | None = Field(default=None, exclude=True, repr=False)
    api_key_env: str | None = None
    api_base: str | None = None
    temperature: float = 1.0
    max_tokens: int = 400
    role_instructions: str = ""
    persona_id: str | None = None
    ideology: str | None = None
    age: int | None = None
    race_ethnicity: str | None = None
    gender: str | None = None
    education: str | None = None
    party: str | None = None
    openness: float | None = None
    conscientiousness: float | None = None
    extraversion: float | None = None
    agreeableness: float | None = None
    neuroticism: float | None = None

    @model_validator(mode="before")
    @classmethod
    def normalize(cls, value):
        data = _persona_fields(_mapping(value, "agent"), complete=False)
        data["model"] = normalize_text(data.get("model"), "model")
        route, _, model_name = data["model"].partition("/")
        if data["model"] != "mock" and (route not in {"openai", "anthropic", "gemini", "ollama"} or not model_name.strip()):
            raise ConfigError("model: use mock, openai/<model>, anthropic/<model>, gemini/<model>, or ollama/<model>.")
        for key in ("display_name", "api_key_env", "api_base"):
            if key in data:
                data[key] = normalize_text(data[key], key, required=False)
        if data.get("display_name") and (":" in data["display_name"] or any(c in data["display_name"] for c in "\n\r\u0085\u2028\u2029")):
            raise ConfigError("display_name: use a single line without a colon.")
        if data.get("api_base"):
            try:
                url = urlsplit(data["api_base"])
                bad_key = any(k.lower().replace("-", "_") in {"key", "api_key", "apikey", "token", "access_token", "authorization", "password", "secret"} for k, _ in parse_qsl(url.query))
                if url.username is not None or url.password is not None or bad_key:
                    raise ConfigError("api_base: embedded credentials are not allowed; use api_key or api_key_env.")
                if url.scheme not in {"http", "https"} or not url.netloc:
                    raise ConfigError("api_base: enter a complete http:// or https:// endpoint URL.")
                _ = url.port
            except ValueError:
                raise ConfigError("api_base: enter a valid endpoint URL without embedded credentials.") from None
        if "api_key" in data:
            secret = data["api_key"]
            if isinstance(secret, SecretStr):
                secret = secret.get_secret_value()
            if is_blank(secret):
                data["api_key"] = None
            elif isinstance(secret, str):
                data["api_key"] = SecretStr(secret)
            else:
                raise ConfigError("api_key: enter a text key.")
        if "role_instructions" in data:
            data["role_instructions"] = normalize_text(data["role_instructions"], "role_instructions", required=False, trim=False) or ""
        if "temperature" in data:
            data["temperature"] = normalize_number(data["temperature"], "temperature", minimum=0)
        if "max_tokens" in data:
            data["max_tokens"] = normalize_number(data["max_tokens"], "max_tokens", integer=True, minimum=1)
        return data

    def persona(self) -> Persona:
        return Persona.model_validate({key: getattr(self, key) for key in PERSONA_COLUMNS})


class Item(FrozenModel):
    item_id: str
    scale: str
    text: str
    type: Literal["likert", "text"]
    min: int | None = None
    max: int | None = None
    min_label: str | None = None
    max_label: str | None = None
    pre: bool = False
    turn: bool = False
    post: bool = False

    @model_validator(mode="before")
    @classmethod
    def normalize(cls, value):
        data = _mapping(value, "item")
        data["item_id"] = normalize_id(data.get("item_id"), "item_id")
        if not re.fullmatch(r"[A-Za-z0-9_]+", data["item_id"]) or data["item_id"] == "reaction":
            raise ConfigError("item_id: use letters, digits and underscores; reaction is reserved.")
        for key in ("scale", "text", "type"):
            data[key] = normalize_text(data.get(key), key, trim=key != "text")
        for key in ("min", "max"):
            if key in data:
                data[key] = normalize_number(data[key], key, integer=True, optional=True)
        for key in ("min_label", "max_label"):
            if key in data:
                data[key] = normalize_text(data[key], key, required=False, trim=False)
        for key in PHASES:
            if key in data:
                data[key] = normalize_boolean(data[key], key)
        return data

    @model_validator(mode="after")
    def validate_bounds(self):
        if self.type == "likert":
            if self.min is None or self.max is None or self.min >= self.max:
                raise ConfigError("likert: min and max must be whole numbers with min < max.")
            if not self.min_label or not self.max_label:
                raise ConfigError("likert: enter non-empty min_label and max_label.")
        elif any(getattr(self, key) is not None for key in ("min", "max", "min_label", "max_label")):
            raise ConfigError("text item: leave min, max, min_label and max_label blank/null.")
        return self


class RunSettings(FrozenModel):
    n_runs: int = 1
    max_turns: int = 15
    speaker_order: tuple[int, ...] = ()  # resolved to actual seat order by Study
    shuffle_order_each_run: bool = False
    allow_pass: bool = True
    silence_visibility: Literal["public", "hidden"] = "public"
    measure_every_n_turns: int = 1
    allow_early_end: bool = False
    remember_pre_survey: bool = False
    randomize_item_order: bool = True
    inner_reaction: bool = False
    private_state: bool = True
    state_max_tokens: int = 1200
    survey_max_tokens: int = 1000
    seed: int | None = None

    @model_validator(mode="before")
    @classmethod
    def normalize(cls, value):
        data = _mapping(value, "run")
        for key in ("n_runs", "max_turns", "measure_every_n_turns", "state_max_tokens", "survey_max_tokens"):
            if key in data:
                data[key] = normalize_number(data[key], key, integer=True, minimum=1)
        if "seed" in data:
            data["seed"] = normalize_number(data["seed"], "seed", integer=True, optional=True)
        for key in ("shuffle_order_each_run", "allow_pass", "allow_early_end", "remember_pre_survey", "randomize_item_order", "inner_reaction", "private_state"):
            if key in data:
                data[key] = normalize_boolean(data[key], key)
        if "speaker_order" in data:
            if not isinstance(data["speaker_order"], (list, tuple)):
                raise ConfigError("speaker_order: enter a list of seat numbers.")
            data["speaker_order"] = tuple(normalize_number(v, "speaker_order", integer=True) for v in data["speaker_order"])
        return data


class Study(FrozenModel):
    study_name: str
    task_instructions: str
    agents: tuple[AgentConfig, ...]
    items: tuple[Item, ...] = ()
    run: RunSettings
    persona_source: Literal["custom", "builtin"] = "custom"
    persona_mode: Literal["fixed", "pool"] = "fixed"
    persona_pool: tuple[Persona, ...] = ()
    survey_instructions: Mapping[str, str] = Field(default_factory=lambda: MappingProxyType(dict(SURVEY_INSTRUCTIONS)))
    prompt_templates: Mapping[str, str] = Field(default_factory=lambda: MappingProxyType({}))

    @field_serializer("survey_instructions", "prompt_templates")
    def serialize_mapping(self, value):
        return dict(value)

    @model_validator(mode="before")
    @classmethod
    def normalize(cls, value):
        data = _mapping(value, "study")
        for key in ("study_name", "task_instructions"):
            data[key] = normalize_text(data.get(key), key, trim=key != "task_instructions")
        for key in ("agents", "items", "persona_pool"):
            rows = data.get(key, () if key != "agents" else None)
            if not isinstance(rows, (tuple, list)):
                raise ConfigError(f"{key}: expected a list of rows.")
            data[key] = tuple(rows)
        for seat, row in enumerate(data["agents"], 1):
            if not isinstance(row, Mapping):
                continue
            if data.get("persona_source", "custom") == "builtin" and set(row) & set(PERSONA_FIELDS):
                raise ConfigError(f"seat {seat}: builtin seats must omit all inline persona fields.")
            if data.get("persona_mode", "fixed") == "pool" and set(row) & set(PERSONA_COLUMNS):
                raise ConfigError(f"seat {seat}: pool seats must omit fixed persona fields and persona_id.")
        run = data.get("run", {})
        run = run.model_dump() if isinstance(run, RunSettings) else _mapping(run, "run")
        if "speaker_order" not in run:
            run["speaker_order"] = tuple(range(1, len(data["agents"]) + 1))
        data["run"] = run
        for key, allowed, defaults in (("survey_instructions", PHASES, SURVEY_INSTRUCTIONS), ("prompt_templates", TEMPLATE_KEYS, {})):
            override = _mapping(data.get(key, {}), key)
            if set(override) - set(allowed):
                raise ConfigError(f"{key}: unknown key; use only registered {'phases' if key == 'survey_instructions' else 'template names'}.")
            for name in override:
                # Empty strings remain valid for no-placeholder helpers.
                if not isinstance(override[name], str):
                    raise ConfigError(f"{key}.{name}: enter text.")
            data[key] = {**defaults, **override}
        return data

    @model_validator(mode="after")
    def validate_study(self):
        count = len(self.agents)
        if count not in (2, 3):
            raise ConfigError("agents: choose exactly 2 or 3 seats.")
        names = [a.display_name for a in self.agents if a.display_name]
        if len(names) != len(set(names)):
            raise ConfigError("display_name: typed names must be unique within the study.")
        if len(self.run.speaker_order) != count or set(self.run.speaker_order) != set(range(1, count + 1)):
            raise ConfigError("speaker_order: include every actual seat exactly once.")
        ids = [i.item_id for i in self.items]
        if len(ids) != len(set(ids)):
            raise ConfigError("items: item_id values must be unique.")
        if self.persona_source == "builtin" and self.persona_pool:
            raise ConfigError("persona_pool: choose custom to supply persona_pool rows.")
        for seat, agent in enumerate(self.agents, 1):
            inline = any(getattr(agent, key) is not None for key in PERSONA_FIELDS)
            if self.persona_source == "builtin" and inline:
                raise ConfigError(f"seat {seat}: builtin seats must omit inline persona fields; choose a persona_id in fixed mode.")
            if self.persona_mode == "pool" and (inline or agent.persona_id is not None):
                raise ConfigError(f"seat {seat}: pool seats must omit fixed persona fields and persona_id.")
            if self.persona_source == "custom" and self.persona_mode == "fixed":
                try:
                    agent.persona()
                except ValidationError as error:
                    raise ConfigError(f"seat {seat}: {_validation_message(error)}") from None
            if self.persona_source == "builtin" and self.persona_mode == "fixed" and not agent.persona_id:
                raise ConfigError(f"seat {seat}: choose a quoted builtin persona_id.")
        if self.persona_mode == "fixed" and self.persona_pool:
            raise ConfigError("persona_pool: fixed studies omit the inactive pool; switch to pool mode to use it.")
        pool = self.persona_pool
        if self.persona_source == "builtin":
            pool = load_builtin_personas()
            if self.persona_mode == "fixed":
                known = {p.persona_id for p in pool}
                used = set()
                for seat, agent in enumerate(self.agents, 1):
                    if agent.persona_id not in known:
                        raise ConfigError(f"seat {seat}: persona_id is not present in the built-in dataset.")
                    if agent.persona_id in used:
                        raise ConfigError(f"seat {seat}: choose a distinct builtin persona_id; this ID is already assigned.")
                    used.add(agent.persona_id)
        if self.persona_mode == "pool":
            validate_pool(pool, self.run.n_runs, count)
        from .prompts import resolve_templates
        resolve_templates(self.prompt_templates)
        # Freeze nested mappings too, since Pydantic frozen alone is shallow.
        object.__setattr__(self, "survey_instructions", MappingProxyType(dict(self.survey_instructions)))
        object.__setattr__(self, "prompt_templates", MappingProxyType(dict(self.prompt_templates)))
        return self

    def warnings(self) -> tuple[str, ...]:
        return tuple(f"Item {item.item_id} is selected for no phase and will not be administered."
                     for item in self.items if not any(getattr(item, phase) for phase in PHASES))


def _validation_message(error: ValidationError) -> str:
    parts = []
    for issue in error.errors(include_input=False, include_url=False, include_context=False):
        path = []
        for component in issue["loc"]:
            path.append(f"row/seat {component + 1}" if isinstance(component, int) else str(component))
        message = "unknown field; remove it" if issue["type"] == "extra_forbidden" else issue["msg"].removeprefix("Value error, ")
        parts.append(f"{'.'.join(path) or 'study'}: {message}")
    return "; ".join(parts)


def validate_study(data: Mapping[str, Any] | Study, *, yaml_input: bool = False) -> Study:
    if isinstance(data, Study):
        return data
    candidate = copy.deepcopy(dict(data))
    if yaml_input:
        for section in ("agents", "persona_pool"):
            entries = candidate.get(section, ())
            if not isinstance(entries, (list, tuple)):
                continue
            for row, entry in enumerate(entries, 1):
                if isinstance(entry, dict) and "persona_id" in entry:
                    try:
                        normalize_id(entry["persona_id"], "persona_id", yaml_input=True)
                    except ConfigError as error:
                        raise ConfigError(f"{section} row/seat {row}: {error}") from None
        # Whole-study import clears all explicit runtime keys, never merges.
        agents = candidate.get("agents", ())
        for agent in agents if isinstance(agents, (list, tuple)) else ():
            if isinstance(agent, dict):
                agent.pop("api_key", None)
    try:
        return Study.model_validate(candidate)
    except DatasetError:
        raise
    except ValidationError as error:
        for issue in error.errors(include_input=False, include_url=False):
            underlying = issue.get("ctx", {}).get("error")
            if isinstance(underlying, DatasetError):
                raise underlying from None
        raise ConfigError(_validation_message(error)) from None


def validate_pool(pool: tuple[Persona, ...], n_runs: int, n_agents: int) -> None:
    ids = [p.persona_id for p in pool]
    if any(not identifier for identifier in ids) or len(ids) != len(set(ids)):
        raise ConfigError("persona_pool: persona_id is required and must be unique on every row.")
    maximum = len(pool) // n_agents
    if n_runs > maximum:
        raise ConfigError(f"persona_pool: {len(pool)} rows allow at most {maximum:,} runs with {n_agents} agents (n_runs <= {maximum}).")


class _UniqueLoader(yaml.SafeLoader):
    pass


def _unique_mapping(loader, node, deep=False):
    loader.flatten_mapping(node)
    result = {}
    for key_node, value_node in node.value:
        key = loader.construct_object(key_node, deep=deep)
        if not isinstance(key, str):
            raise ConfigError(f"YAML line {key_node.start_mark.line + 1}: field names must be strings.")
        if key in result:
            # Do not echo source keys, which may contain sensitive text.
            raise ConfigError(f"YAML line {key_node.start_mark.line + 1}: duplicate key; keep one value per field.")
        if key == "item_id" and isinstance(value_node, yaml.ScalarNode) and value_node.tag == "tag:yaml.org,2002:int":
            result[key] = value_node.value  # retain digits/leading zeros in source
        else:
            result[key] = loader.construct_object(value_node, deep=deep)
    return result


_UniqueLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _unique_mapping)


def parse_study_yaml(text: str) -> dict:
    """Read a candidate with the same safe, duplicate-rejecting YAML loader."""
    try:
        data = yaml.load(text, Loader=_UniqueLoader)
    except yaml.YAMLError as error:
        mark = getattr(error, "problem_mark", None)
        location = f" at line {mark.line + 1}" if mark else ""
        raise ConfigError(f"Invalid YAML{location}; check indentation and field syntax.") from None
    if not isinstance(data, dict):
        raise ConfigError("Study YAML must contain one object with study fields.")
    return data


def load_study_yaml(text: str) -> Study:
    return validate_study(parse_study_yaml(text), yaml_input=True)


def read_study_data(path: str | Path) -> dict:
    try:
        text = Path(path).read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError):
        raise ConfigError("Cannot read the study file as UTF-8 YAML.") from None
    return parse_study_yaml(text)


def load_study(path: str | Path) -> Study:
    return validate_study(read_study_data(path), yaml_input=True)


def with_custom_persona_pool(data: Mapping[str, Any], personas: tuple[Persona, ...]) -> dict:
    """Apply the CLI's explicit pool override before validating active fields.

    Preserve seat configuration and unknown fields for strict validation. Only
    the now-inactive persona fields/IDs are removed; the caller's data is intact.
    """
    candidate = copy.deepcopy(dict(data))
    candidate.update(persona_source="custom", persona_mode="pool",
                     persona_pool=[persona.model_dump() for persona in personas])
    entries = candidate.get("agents", ())
    for agent in entries if isinstance(entries, (list, tuple)) else ():
        if isinstance(agent, dict):
            for key in (*PERSONA_FIELDS, "persona_id"):
                agent.pop(key, None)
    return candidate


def study_to_dict(study: Study) -> dict:
    # Deliberately serialize from approved config fields, never execution states
    # or resolved credentials. Reference-shaped builtin YAML stays reference-shaped.
    agents = []
    for agent in study.agents:
        row = agent.model_dump(exclude_none=True)
        if study.persona_mode == "pool" or study.persona_source == "builtin":
            for key in PERSONA_FIELDS:
                row.pop(key, None)
        if study.persona_mode == "pool":
            row.pop("persona_id", None)
        agents.append(row)
    data = {
        "study_name": study.study_name, "task_instructions": study.task_instructions,
        "persona_source": study.persona_source, "persona_mode": study.persona_mode,
        "agents": agents, "run": study.run.model_dump(mode="json"),
        "items": [item.model_dump(exclude_none=True) for item in study.items],
        "survey_instructions": dict(study.survey_instructions),
        "prompt_templates": dict(study.prompt_templates),
    }
    if study.persona_source == "custom" and study.persona_mode == "pool":
        data["persona_pool"] = [persona.model_dump(exclude_none=True) for persona in study.persona_pool]
    return data


class _QuotedID(str):
    pass


class _StudyDumper(yaml.SafeDumper):
    pass


_StudyDumper.add_representer(_QuotedID, lambda dumper, value: dumper.represent_scalar("tag:yaml.org,2002:str", value, style='"'))


def dump_study_yaml(study: Study) -> str:
    data = study_to_dict(study)
    for section in ("agents", "persona_pool"):
        for row in data.get(section, ()):
            if "persona_id" in row:
                row["persona_id"] = _QuotedID(row["persona_id"])
    return yaml.dump(data, Dumper=_StudyDumper, sort_keys=False, allow_unicode=True)


def save_study(study: Study, path: str | Path) -> None:
    Path(path).write_text(dump_study_yaml(study), encoding="utf-8")


def _check_excel_nonfinite(data: bytes, worksheet_path: str, header: list, filename: str) -> None:
    """Name the cell for non-finite XML numbers before openpyxl rejects them.

    Excel cannot represent these as normal numeric cells. Checking the original
    cell prevents an unreadable value becoming a missing cell or a generic error.
    The pinned openpyxl reader supplies the actual first worksheet's ZIP path.
    """
    from zipfile import ZipFile
    from openpyxl.utils.cell import column_index_from_string
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(io.BytesIO(data)) as archive:
        sheet_xml = fromstring(archive.read(worksheet_path))
    for cell in sheet_xml.iter(namespace + "c"):
        if cell.get("t", "n") != "n":
            continue
        node = cell.find(namespace + "v")
        if node is None or not node.text:
            continue
        try:
            nonfinite = not Decimal(node.text).is_finite()
        except InvalidOperation:
            continue  # other malformed cells receive the ordinary file error
        if nonfinite:
            coordinate = re.fullmatch(r"([A-Z]+)([0-9]+)", cell.get("r", ""))
            if coordinate:
                column = column_index_from_string(coordinate[1]) - 1
                field = header[column] if column < len(header) else coordinate[1]
                raise ConfigError(f"{Path(filename).name} row {coordinate[2]}: {field}: NaN and infinity are not allowed.")
            raise ConfigError(f"{Path(filename).name}: NaN and infinity are not allowed in numeric cells.")


def _read_table(data: bytes, filename: str, columns: tuple[str, ...], required: set[str]) -> list[tuple[int, dict]]:
    suffix = Path(filename).suffix.lower()
    rows = []
    try:
        if suffix == ".csv":
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig"), newline=""), strict=True)
            header = next(reader, [])
            for values in reader:
                rows.append((reader.line_num, values))
        elif suffix == ".xlsx":
            # Read original cell types; pandas' automatic NA/type coercion would
            # lose NULL strings, duplicate headers, and numeric ID distinctions.
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                sheet = workbook.worksheets[0]
                iterator = sheet.iter_rows(values_only=True)
                header = list(next(iterator, []))
                _check_excel_nonfinite(data, sheet._worksheet_path, header, filename)
                rows = [(index, list(values)) for index, values in enumerate(iterator, 2)]
            finally:
                workbook.close()
        else:
            raise ConfigError("Use a .csv or .xlsx file; Excel imports read the first worksheet.")
    except (UnicodeError, csv.Error, OSError, ValueError, KeyError, IndexError, BadZipFile, ParseError) as error:
        if isinstance(error, ConfigError):
            raise
        raise ConfigError(f"{Path(filename).name}: cannot read the file; use a valid UTF-8 CSV or .xlsx workbook.") from None
    if not header or any(not isinstance(key, str) or not key for key in header):
        raise ConfigError(f"{Path(filename).name}: every column must have an exact documented header.")
    if len(header) != len(set(header)):
        raise ConfigError(f"{Path(filename).name}: duplicate column headers are not allowed.")
    if set(header) - set(columns):
        raise ConfigError(f"{Path(filename).name}: unknown column; use only the documented headers.")
    missing = required - set(header)
    if missing:
        raise ConfigError(f"{Path(filename).name}: missing required columns: {', '.join(sorted(missing))}.")
    result = []
    for row_number, values in rows:
        if all(is_blank(v) for v in values):
            continue
        if len(values) != len(header):
            raise ConfigError(f"{Path(filename).name} row {row_number}: column count does not match the header.")
        result.append((row_number, dict(zip(header, values))))
    return result


def import_personas(data: bytes, filename: str, *, mode: Literal["fixed", "pool"], n_agents: int | None = None) -> tuple[Persona, ...]:
    if mode not in {"fixed", "pool"}:
        raise ConfigError("persona_mode: choose fixed or pool.")
    required = set(DEMOGRAPHIC_FIELDS) | ({"persona_id"} if mode == "pool" else set())
    rows = _read_table(data, filename, PERSONA_COLUMNS, required)
    if mode == "fixed" and len(rows) != n_agents:
        raise ConfigError(f"{Path(filename).name}: fixed upload needs exactly {n_agents} rows, one per selected seat; found {len(rows)}.")
    result, seen = [], set()
    for row_number, row in rows:
        try:
            persona = Persona.model_validate(row)
            if mode == "pool":
                if not persona.persona_id:
                    raise ConfigError("persona_id: enter a non-empty ID.")
                if persona.persona_id in seen:
                    raise ConfigError("persona_id: duplicate ID; pool IDs must be unique.")
                seen.add(persona.persona_id)
            result.append(persona)
        except (ValidationError, ConfigError) as error:
            message = _validation_message(error) if isinstance(error, ValidationError) else str(error)
            raise ConfigError(f"{Path(filename).name} row {row_number}: {message}") from None
    return tuple(result)


def import_items(data: bytes, filename: str) -> tuple[Item, ...]:
    rows = _read_table(data, filename, ITEM_COLUMNS, set(ITEM_COLUMNS))
    result, seen = [], set()
    for row_number, row in rows:
        try:
            for phase in PHASES:
                row[phase] = normalize_boolean(row[phase], phase, table=True)
            item = Item.model_validate(row)
            if item.item_id in seen:
                raise ConfigError("item_id: duplicate ID; questionnaire IDs must be unique.")
            seen.add(item.item_id)
            result.append(item)
        except (ValidationError, ConfigError) as error:
            message = _validation_message(error) if isinstance(error, ValidationError) else str(error)
            raise ConfigError(f"{Path(filename).name} row {row_number}: {message}") from None
    return tuple(result)


def read_import(path: str | Path) -> bytes:
    try:
        return Path(path).read_bytes()
    except OSError:
        raise ConfigError("Cannot read the selected import file.") from None


@lru_cache(maxsize=1)
def load_builtin_personas() -> tuple[Persona, ...]:
    try:
        data = resources.files("convolab").joinpath("data", BUILTIN_DATASET_FILE).read_bytes()
    except (OSError, FileNotFoundError):
        raise DatasetError(f"The built-in file {BUILTIN_DATASET_FILE} is missing or unreadable. Reinstall the application data before starting.") from None
    if hashlib.sha256(data).hexdigest() != BUILTIN_DATASET_SHA256:
        raise DatasetError(f"The built-in file {BUILTIN_DATASET_FILE} has changed and failed its SHA-256 check. Restore the packaged file before starting.")
    try:
        rows = import_personas(data, BUILTIN_DATASET_FILE, mode="pool")
        if len(rows) != BUILTIN_DATASET_ROWS:
            raise ConfigError(f"Expected {BUILTIN_DATASET_ROWS} rows.")
    except ConfigError:
        raise DatasetError(f"The built-in file {BUILTIN_DATASET_FILE} is invalid or has the wrong row count. Restore the packaged file before starting.") from None
    return rows


@dataclass(frozen=True)
class FrozenExecution:
    study: Study
    fixed_personas: tuple[Persona, ...]
    persona_pool: tuple[Persona, ...]
    persona_dataset: Mapping[str, Any]


def freeze_execution(study: Study) -> FrozenExecution:
    # Revalidate a detached, active-only serialization; never deepcopy proxies.
    frozen = validate_study(study_to_dict(study))
    fixed, pool = (), ()
    metadata = {"source": study.persona_source}
    if frozen.persona_source == "builtin":
        dataset = load_builtin_personas()
        metadata.update(dataset_id=BUILTIN_DATASET_ID, file=BUILTIN_DATASET_FILE,
                        n_rows=BUILTIN_DATASET_ROWS, sha256=BUILTIN_DATASET_SHA256)
        if frozen.persona_mode == "fixed":
            by_id = {p.persona_id: p for p in dataset}
            fixed = tuple(Persona.model_validate(by_id[a.persona_id].model_dump()) for a in frozen.agents)
        else:
            pool = tuple(Persona.model_validate(p.model_dump()) for p in dataset)
    elif frozen.persona_mode == "fixed":
        fixed = tuple(a.persona() for a in frozen.agents)
    else:
        pool = tuple(Persona.model_validate(p.model_dump()) for p in frozen.persona_pool)
    return FrozenExecution(frozen, fixed, pool, MappingProxyType(metadata))


def resolve_credentials(study: Study, environment: Mapping[str, str] | None = None) -> tuple[SecretStr | None, ...]:
    environment = os.environ if environment is None else environment
    result = []
    for seat, agent in enumerate(study.agents, 1):
        explicit = agent.api_key.get_secret_value() if agent.api_key else None
        if explicit and explicit.strip():
            result.append(SecretStr(explicit))
            continue
        if agent.model == "mock" or agent.model.startswith("ollama/"):
            result.append(None)
            continue
        key = environment.get(agent.api_key_env, "") if agent.api_key_env else ""
        if not isinstance(key, str) or not key.strip():
            raise ConfigError(f"seat {seat}: supply a non-empty API key or a resolving api_key_env reference before starting.")
        result.append(SecretStr(key))
    return tuple(result)


def apply_persona_import(study: Study, data: bytes, filename: str) -> Study:
    if study.persona_source != "custom":
        raise ConfigError("Persona uploads require persona_source custom.")
    rows = import_personas(data, filename, mode=study.persona_mode, n_agents=len(study.agents))
    candidate = study_to_dict(study)
    # Runtime secrets survive persona-only edits, but never enter serialization.
    for agent, old in zip(candidate["agents"], study.agents):
        agent["api_key"] = old.api_key
    if study.persona_mode == "pool":
        candidate["persona_pool"] = [p.model_dump() for p in rows]
    else:
        for agent, persona in zip(candidate["agents"], rows):
            agent.update(persona.model_dump())
    return validate_study(candidate)


def apply_item_import(study: Study, data: bytes, filename: str) -> Study:
    items = import_items(data, filename)
    candidate = study_to_dict(study)
    for agent, old in zip(candidate["agents"], study.agents):
        agent["api_key"] = old.api_key
    candidate["items"] = [item.model_dump() for item in items]
    return validate_study(candidate)


def active_editor_data(editor: Mapping[str, Any]) -> dict:
    """Detach active fields without copying personas across source or mode.

    This returns a candidate, not a validated study. Empty selections after a
    mode/source change must be filled and validated before execution.
    """
    data = copy.deepcopy(dict(editor))
    source, mode = data.get("persona_source", "custom"), data.get("persona_mode", "fixed")
    if mode != "pool" or source != "custom":
        data.pop("persona_pool", None)
    for agent in data.get("agents", ()):
        if mode == "pool" or source == "builtin":
            for key in PERSONA_FIELDS:
                agent.pop(key, None)
        if mode == "pool":
            agent.pop("persona_id", None)
    return data


def resize_editor_seats(editor: Mapping[str, Any], count: int) -> dict:
    count = normalize_number(count, "agent count", integer=True)
    if count not in (2, 3):
        raise ConfigError("agent count: choose 2 or 3.")
    data = copy.deepcopy(dict(editor))
    agents = list(data.get("agents", ()))[:count]
    agents.extend({} for _ in range(count - len(agents)))
    data["agents"] = agents
    data.setdefault("run", {})["speaker_order"] = list(range(1, count + 1))
    return data


@dataclass
class StudyEditor:
    """Source-specific editor buffers, separate from the frozen study.

    Switching controls selects a buffer; it never converts or copies personas
    between buffers. An unfilled new selection produces an invalid candidate
    until the researcher supplies the required data. UI wiring is Stage 6.
    """
    settings: dict
    custom_fixed: list[dict]
    builtin_fixed_ids: list[str | None]
    custom_pool: list[dict]
    persona_source: Literal["custom", "builtin"]
    persona_mode: Literal["fixed", "pool"]

    @classmethod
    def from_study(cls, study: Study) -> StudyEditor:
        data = study_to_dict(study)
        fixed = [{} for _ in study.agents]
        builtin_ids = [None for _ in study.agents]
        for index, (row, agent) in enumerate(zip(data["agents"], study.agents)):
            if study.persona_mode == "fixed":
                if study.persona_source == "custom":
                    fixed[index] = {key: row[key] for key in PERSONA_COLUMNS if key in row}
                else:
                    builtin_ids[index] = agent.persona_id
            for key in PERSONA_COLUMNS:
                row.pop(key, None)
            row["api_key"] = agent.api_key
        pool = data.pop("persona_pool", [])
        data.pop("persona_source")
        data.pop("persona_mode")
        return cls(data, fixed, builtin_ids, pool, study.persona_source, study.persona_mode)

    def candidate(self) -> dict:
        data = copy.deepcopy(self.settings)
        data.update(persona_source=self.persona_source, persona_mode=self.persona_mode)
        if self.persona_mode == "fixed":
            for index, row in enumerate(data["agents"]):
                if self.persona_source == "custom":
                    row.update(copy.deepcopy(self.custom_fixed[index]))
                elif self.builtin_fixed_ids[index] is not None:
                    row["persona_id"] = self.builtin_fixed_ids[index]
        elif self.persona_source == "custom":
            data["persona_pool"] = copy.deepcopy(self.custom_pool)
        return data

    def validated(self) -> Study:
        return validate_study(self.candidate())

    def resize(self, count: int) -> None:
        candidate = resize_editor_seats(self.settings, count)
        fixed = copy.deepcopy(self.custom_fixed[:count])
        fixed.extend({} for _ in range(count - len(fixed)))
        identifiers = self.builtin_fixed_ids[:count]
        identifiers.extend(None for _ in range(count - len(identifiers)))
        self.settings, self.custom_fixed, self.builtin_fixed_ids = candidate, fixed, identifiers
